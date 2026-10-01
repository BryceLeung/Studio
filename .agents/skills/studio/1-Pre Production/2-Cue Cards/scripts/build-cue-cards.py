"""Build editable Studio cue cards on Windows, macOS, or Linux."""

from argparse import ArgumentParser
from hashlib import sha256
from pathlib import Path
import json
import math
import os
import re

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt


SLIDE_WIDTH = 960
SLIDE_HEIGHT = 540
FONT_NAME = "Arial"
CUE_FONT_PT = 36
CUE_WIDTH = 830
CUE_HEIGHT = 426
MAX_LINES = 9
BACKGROUND = RGBColor(19, 21, 28)
WHITE = RGBColor(255, 255, 255)


def metric_font(explicit: Path | None) -> ImageFont.FreeTypeFont:
    if explicit is not None:
        return ImageFont.truetype(str(explicit), CUE_FONT_PT)
    candidates = []
    if os.environ.get("WINDIR"):
        candidates.append(str(Path(os.environ["WINDIR"]) / "Fonts" / "arial.ttf"))
    candidates.extend(
        [
            "/Library/Fonts/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Microsoft/Arial.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
            "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            "Arial.ttf",
            "LiberationSans-Regular.ttf",
            "DejaVuSans.ttf",
        ]
    )
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, CUE_FONT_PT)
        except OSError:
            pass
    raise RuntimeError("No metric font found; pass --font-file with an Arial-compatible TTF font")


def visual_line_count(text: str, font: ImageFont.FreeTypeFont) -> int:
    lines = 0
    for hard_line in text.split("\n"):
        if not hard_line:
            lines += 1
            continue
        current = ""
        for word in hard_line.split():
            if font.getlength(word) > CUE_WIDTH:
                raise ValueError(f"A word is wider than the cue box: {word}")
            candidate = f"{current} {word}".strip()
            if current and font.getlength(candidate) > CUE_WIDTH:
                lines += 1
                current = word
            else:
                current = candidate
        lines += 1
    return lines


def fits(text: str, font: ImageFont.FreeTypeFont) -> bool:
    ascent, descent = font.getmetrics()
    height_limit = max(1, math.floor((CUE_HEIGHT - 12) / (ascent + descent)))
    return visual_line_count(text, font) <= min(MAX_LINES, height_limit)


def split_overflow(text: str, font: ImageFont.FreeTypeFont) -> list[str]:
    if fits(text, font):
        return [text]
    lines = text.split("\n")
    if len(lines) > 1:
        middle = math.ceil(len(lines) / 2)
        first = "\n".join(lines[:middle]).strip()
        second = "\n".join(lines[middle:]).strip()
    else:
        sentences = re.split(r"(?<=[.!?…])\s+(?=[A-Z“‘])", text)
        if len(sentences) > 1:
            middle = math.ceil(len(sentences) / 2)
            first = " ".join(sentences[:middle])
            second = " ".join(sentences[middle:])
        else:
            words = text.split()
            if len(words) < 2:
                raise ValueError("Cue text cannot be split without changing a word")
            middle = math.ceil(len(words) / 2)
            first = " ".join(words[:middle])
            second = " ".join(words[middle:])
    if not first or not second:
        raise ValueError("Cue text cannot be split into two nonempty slides")
    return split_overflow(first, font) + split_overflow(second, font)


def add_background(slide) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Pt(SLIDE_WIDTH), Pt(SLIDE_HEIGHT))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BACKGROUND
    shape.line.fill.background()


def add_text(slide, text: str, x: int, y: int, width: int, height: int, size: int,
             align: PP_ALIGN = PP_ALIGN.LEFT, bold: bool = False) -> None:
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = WHITE


def build_deck(plan: dict, font: ImageFont.FreeTypeFont, output: Path) -> int:
    presentation = Presentation()
    presentation.slide_width = Pt(SLIDE_WIDTH)
    presentation.slide_height = Pt(SLIDE_HEIGHT)
    blank = presentation.slide_layouts[6]

    title_slide = presentation.slides.add_slide(blank)
    add_background(title_slide)
    add_text(title_slide, plan["title"], 70, 145, 820, 210, 54, PP_ALIGN.CENTER, True)
    add_text(title_slide, "CUE CARDS", 70, 375, 820, 45, 19, PP_ALIGN.CENTER)

    cues = [part for text in plan["slides"] for part in split_overflow(text, font)]
    for number, cue in enumerate(cues, 1):
        slide = presentation.slides.add_slide(blank)
        add_background(slide)
        add_text(slide, cue, 65, 52, CUE_WIDTH, CUE_HEIGHT, CUE_FONT_PT)
        add_text(slide, f"{number:02} / {len(cues):02}", 800, 485, 88, 30, 16, PP_ALIGN.RIGHT)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return len(cues)


def main() -> None:
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("project", type=Path, help="OneDrive Studio project folder")
    parser.add_argument("--output-root", type=Path, help="Directory for generated decks; defaults to the cue-card step folder")
    parser.add_argument("--font-file", type=Path, help="Arial-compatible TTF for measuring text on this computer")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing decks after reviewing them for human edits")
    args = parser.parse_args()

    project = args.project.resolve(strict=True)
    approved = project / "1-Pre Production" / "1-Script Review" / "Approved"
    working = project / "1-Pre Production" / "2-Cue Cards" / "working"
    destination = args.output_root or working.parent
    sources = sorted(approved.glob("*.docx"))
    plans = sorted(working.glob("*.json"))
    if not sources or {path.stem for path in sources} != {path.stem for path in plans}:
        parser.error("Approved Word scripts and cue-card plans must be a nonempty matching set")

    jobs = []
    for plan_file in plans:
        source = approved / f"{plan_file.stem}.docx"
        plan = json.loads(plan_file.read_text(encoding="utf-8"))
        if sha256(source.read_bytes()).hexdigest().upper() != plan["source_sha256"]:
            raise ValueError(f"Approved script changed after planning: {source}")
        output = destination / f"{plan_file.stem}.pptx"
        if output.exists() and not args.overwrite:
            raise FileExistsError(f"Review the existing deck before passing --overwrite: {output}")
        jobs.append((plan, output))

    font = metric_font(args.font_file)
    for plan, output in jobs:
        count = build_deck(plan, font, output)
        print(f"Created {output} with {count + 1} slides ({count} speaking)")


if __name__ == "__main__":
    main()
